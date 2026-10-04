# -*- coding: utf-8 -*-
"""
给《Claude版信安初稿.docx》插图+调格式（2026-09-15）
- 5个占位段落替换为居中图片+图注（宋体10.5pt）
- 4张表格设置列宽适配页面
"""
import io
import os
import sys

import docx
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

DESKTOP = os.path.join(os.path.expanduser("~"), "Desktop")
REPORT = os.path.join(DESKTOP, "Claude版信安初稿.docx")
TMP = os.path.join(os.environ.get("TEMP", r"C:\Users\34612\AppData\Local\Temp"), "oldreport")
SHOTS = r"D:\DeskApp\python_xiangmu\python-sbom-security-tool\temp_shots"

# 图注 -> (图片路径, 宽度cm)
IMAGES = {
    "【此处插入图片：图2-1 系统总体架构图】":
        (os.path.join(TMP, "arch.png"), 14.5, "图2-1 系统总体架构图"),
    "【此处插入图片：图2-2 系统主界面（初始状态）】":
        (os.path.join(SHOTS, "frames", "frame_001.jpg"), 12.0, "图2-2 系统主界面"),
    "【此处插入图片：图2-3 扫描结果总览界面】":
        (os.path.join(SHOTS, "13_scan_result.png"), 13.0, "图2-3 扫描结果总览界面"),
    "【此处插入图片：图2-4 Skill安全检测界面】":
        (os.path.join(SHOTS, "10_skill_detect.png"), 13.0, "图2-4 Skill安全检测界面"),
    "【此处插入图片：图2-5 AI辅助研判界面】":
        (os.path.join(SHOTS, "09_ai_report.png"), 13.0, "图2-5 AI辅助研判界面"),
}

# 表格列宽（按表索引: 0=表3-1, 1=表3-2, 2=表3-3, 3=表3-4）
TABLE_WIDTHS = {
    0: [2.6, 1.8, 2.8, 2.8, 2.2],
    1: [6.2, 1.8, 1.8, 1.8, 1.9],
    2: [4.6, 1.7, 1.5, 1.6, 2.4, 1.5],
    3: [2.4, 1.8, 2.4, 2.0, 2.0, 2.2],
}


def set_font(run, size=10.5, bold=False):
    run.font.name = "宋体"
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


# 统一转成标准PNG（python-docx对部分JPEG变体解析失败）
from PIL import Image
CLEAN = os.path.join(TMP, "clean")
os.makedirs(CLEAN, exist_ok=True)
clean_map = {}
for key, (path, width, caption) in IMAGES.items():
    clean = os.path.join(CLEAN, os.path.basename(path).rsplit(".", 1)[0] + ".png")
    Image.open(path).convert("RGB").save(clean, "PNG")
    clean_map[key] = (clean, width, caption)
IMAGES = clean_map

doc = docx.Document(REPORT)

# ---- 1. 插图 ----
replaced = 0
for p in list(doc.paragraphs):
    text = p.text.strip()
    if text not in IMAGES:
        continue
    path, width, caption = IMAGES[text]
    if not os.path.isfile(path):
        print("缺图:", path)
        continue
    # 清空占位文字
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    # 图片
    r_img = p.add_run()
    r_img.add_picture(path, width=Cm(width))
    # 换行+图注
    r_cap = p.add_run()
    r_cap.add_break()
    set_font(r_cap, 10.5, False)
    r_cap.add_text(caption)
    replaced += 1
    print("已插图:", caption, "<-", os.path.basename(path))

# ---- 2. 表格列宽 + 居中 ----
for idx, t in enumerate(doc.tables):
    if idx not in TABLE_WIDTHS:
        continue
    widths = TABLE_WIDTHS[idx]
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for row in t.rows:
        for j, cell in enumerate(row.cells):
            cell.width = Cm(widths[j])
            for para in cell.paragraphs:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in para.runs:
                    run.font.name = "宋体"
                    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    print("表%d列宽已设:" % (idx + 1), widths)

doc.save(REPORT)
print("完成,替换占位%d处,已保存:" % replaced, REPORT)
