"""
生成简洁版汇报文档（学生口吻，5个板块）
"""
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
import os

doc = Document()

style = doc.styles['Normal']
style.font.name = 'Microsoft YaHei'
style.font.size = Pt(11.5)
style.paragraph_format.line_spacing = 1.6
style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')

for sec in doc.sections:
    sec.top_margin = Cm(2.5)
    sec.bottom_margin = Cm(2.5)
    sec.left_margin = Cm(2.8)
    sec.right_margin = Cm(2.8)

# ===== 封面 =====
for _ in range(7):
    doc.add_paragraph('')

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('Python 依赖漏洞检测工具')
r.font.size = Pt(26)
r.font.bold = True
r.font.color.rgb = RGBColor(0x2c, 0x3e, 0x50)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('暑期学习进展汇报')
r.font.size = Pt(14)
r.font.color.rgb = RGBColor(0x7f, 0x8c, 0x8d)

doc.add_paragraph('')
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('2026年7月')
r.font.size = Pt(11)
r.font.color.rgb = RGBColor(0x95, 0xa5, 0xa6)

doc.add_page_break()

# ===== 一、做了什么 =====
doc.add_heading('一、这个月做了什么', level=1)

doc.add_paragraph(
    '做了一个 Python 项目的依赖漏洞扫描工具。简单说就是：上传一个 requirements.txt，'
    '工具会告诉你项目用了哪些第三方库、这些库有没有已知的安全漏洞、有多严重、应该先修哪个。'
)

doc.add_paragraph(
    '核心思路是"多查几个数据库，交叉验证"。同一个漏洞，我会去 OSV、NVD、GitHub Advisory '
    '三个数据库都查一遍，三方都说有问题才标"高置信度"，只有一方说的标"低置信度"。'
    '这样就不会被单一数据库的错误数据误导。'
)

doc.add_paragraph(
    '还做了一个简单的可视化界面，扫描完自动出几张图表，能直观看到漏洞分布情况。'
)

# ===== 二、流程图 =====
doc.add_heading('二、工具是怎么跑的', level=1)

diagram_path = r'C:\Users\34612\Desktop\项目流程图.png'
if os.path.exists(diagram_path):
    doc.add_picture(diagram_path, width=Inches(5.6))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph(
    '流程很简单，五步走：上传文件 → 解析出用了哪些库 → 去三个数据库查漏洞 → '
    '交叉验证给置信度 → 生成报告和图表。'
)

# ===== 三、和上次比改了哪里 =====
doc.add_heading('三、和上个版本比改了哪里', level=1)

doc.add_paragraph(
    '上次老师反馈说功能太散、堆了很多东西但都不深入。这个月主要做了三件事：'
)

changes = [
    '砍掉了外围功能（国密签名、依赖图谱、多生态支持等），只保留最核心的漏洞检测',
    '把漏洞检测做深了——原来是"查到就展示"，现在是"多源交叉验证 → 置信度评级 → 利用概率预测 → 修复优先级排序"',
    '加了可视化图表（风险分布、风险矩阵散点图等5张图），扫描完直接看到整体安全状况',
]
for c in changes:
    doc.add_paragraph(c, style='List Bullet')

doc.add_paragraph(
    '功能数量比上个版本少了，但核心功能的深度提高了。每个漏洞结论都有 NVD 官方链接可以点开验证。'
)

# ===== 四、现在还存在的问题 =====
doc.add_heading('四、目前的问题和下一步计划', level=1)

doc.add_paragraph('目前这个版本还有几个明显的不足：')
problems = [
    '只支持 Python 项目，前端（npm）和 Java（Maven）还扫不了',
    '只能扫依赖的漏洞，还没有做许可证合规检查和供应链攻击检测',
    'API 调用比较慢（每次扫描要联网查好几个数据库），没有做离线缓存',
    '报告目前只有 DOCX 和 PDF，不够丰富',
]
for prob in problems:
    doc.add_paragraph(prob, style='List Bullet')

doc.add_paragraph('')
doc.add_paragraph('接下来 8 月份计划做的事：')
plans = [
    '加上 npm（package.json）的支持，让前端项目也能扫',
    '做扫描历史记录功能，方便对比两次扫描之间的漏洞变化',
    '如果时间够，接入国内漏洞数据库（CNNVD）',
    '持续优化，争取 10 月份能拿去参加省赛',
]
for plan in plans:
    doc.add_paragraph(plan, style='List Bullet')

# ===== 五、参考方向 =====
doc.add_heading('五、参考了哪些方向', level=1)

doc.add_paragraph(
    '看了今年国赛的 300 个作品名单。发现做 AI/大模型安全的特别多（大概 50 多个），'
    '做密码学隐私计算的也不少（40 个左右），但做软件供应链安全方向的只有 8 个左右，'
    '专门做依赖漏洞检测的一个都没有。'
)

doc.add_paragraph(
    '所以觉得这个方向虽然现在做的人少，但反而是个机会。后面可以在这个方向上继续深入，'
    '不跟大多数人挤在 AI 安全那条赛道上。'
)

doc.add_paragraph(
    '最近也在看实验室学长学姐做的项目（反洗钱系统、恶意流量检测），他们的技术名词比较多，'
    '但拆开来看核心逻辑也没有那么复杂。接下来打算花时间把他们的技术方案搞清楚，'
    '能借鉴的地方借鉴过来。'
)

# ===== 保存 =====
out = r'C:\Users\34612\Desktop\项目汇报文档.docx'
try:
    os.remove(out)
except:
    pass
doc.save(out)
print(f'OK: {out}')
