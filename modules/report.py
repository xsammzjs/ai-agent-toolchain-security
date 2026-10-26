"""
报告生成模块（v3.0）
====================
生成 DOCX 审计报告
"""
import os
from datetime import datetime
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn


def _set_cell(table, row_idx, col_idx, text):
    """安全写入表格单元格"""
    table.rows[row_idx].cells[col_idx].text = text


def generate_docx_report(scan_result: dict, sbom_data: dict) -> str:
    """生成 DOCX 格式审计报告。返回文件路径。"""
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Microsoft YaHei'
    style.font.size = Pt(10)
    style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')

    # ===== 标题页 =====
    title = doc.add_heading('Python 依赖安全审计报告', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph(f'生成时间：{datetime.now().strftime("%Y-%m-%d %H:%M")}')
    doc.add_paragraph(f'扫描工具：AI 技能与依赖漏洞检测平台')
    doc.add_paragraph('')

    # ===== 一、扫描摘要 =====
    doc.add_heading('一、扫描摘要', level=1)
    summary = scan_result.get("summary", {})
    dep_list = scan_result.get("dep_list", [])
    vulns = scan_result.get("vulns", [])

    summary_data = [
        ("扫描时间", summary.get("scan_time", "未知")),
        ("依赖总数", str(summary.get("total_dependencies", 0))),
        ("其中间接依赖", str(summary.get("indirect_count", 0))),
        ("漏洞总数", str(summary.get("total_vulnerabilities", 0))),
        ("严重/高危漏洞", str(summary.get("critical_count", 0))),
        ("需立即修复", str(summary.get("priority_1_count", 0))),
        ("高置信度漏洞", str(summary.get("high_confidence_count", 0))),
        ("可修复漏洞", str(summary.get("fixable_count", 0))),
        ("活跃利用风险", str(summary.get("actively_exploited_count", 0))),
    ]

    summary_table = doc.add_table(rows=len(summary_data), cols=2, style='Light Grid Accent 1')
    for i, (key, val) in enumerate(summary_data):
        _set_cell(summary_table, i, 0, key)
        _set_cell(summary_table, i, 1, val)

    doc.add_paragraph('')

    # ===== 二、漏洞清单 =====
    doc.add_heading('二、漏洞清单', level=1)

    if not vulns:
        doc.add_paragraph('未检测到已知漏洞，项目依赖安全状况良好。')
    else:
        doc.add_paragraph(f'共检测到 {len(vulns)} 个已知漏洞，按修复优先级排序：')
        doc.add_paragraph('')

        vuln_table = doc.add_table(rows=1, cols=10, style='Light Grid Accent 1')
        headers = ['优先级', 'CVE编号', '组件', '版本', '层级', '风险', 'CVSS', '置信度', 'EPSS', '修复版本']
        for i, header in enumerate(headers):
            _set_cell(vuln_table, 0, i, header)

        for v in vulns:
            row = vuln_table.add_row()
            priority_label = v.get("fix_priority", {}).get("label", "未知")
            cve_id = v.get("cve_id", "")
            pkg = f"{v.get('package', '')}"
            version = v.get("version", "")
            layer = "间接" if v.get("is_transitive") else "直接"
            severity = v.get("severity", "")
            cvss = str(v.get("cvss_score", 0))
            confidence = f"{v.get('confidence', {}).get('badge', '')} {v.get('confidence', {}).get('level', '')}"
            epss = f"{v.get('epss_score', 0):.1%}" if v.get('epss_score') else "N/A"
            fixed = v.get("fixed_version", "未知")

            cells = row.cells
            cells[0].text = priority_label
            cells[1].text = cve_id
            cells[2].text = pkg
            cells[3].text = version
            cells[4].text = layer
            cells[5].text = severity
            cells[6].text = cvss
            cells[7].text = confidence
            cells[8].text = epss
            cells[9].text = fixed

    doc.add_paragraph('')

    # ===== 三、漏洞详情 =====
    doc.add_heading('三、漏洞详情与证据溯源', level=1)

    for i, v in enumerate(vulns[:20], 1):
        doc.add_heading(f'{i}. {v.get("cve_id", "")} — {v.get("package", "")} {v.get("version", "")}', level=2)

        detail_data = [
            ("风险等级", v.get("severity", "未知")),
            ("CVSS 分数", f"{v.get('cvss_score', 0)} (来源: {v.get('score_source', '未知')})"),
            ("分数类型", "精确" if v.get("score_type") == "exact" else "估算"),
            ("置信度", f"{v.get('confidence', {}).get('badge', '')} {v.get('confidence', {}).get('level', '')} 级 - {v.get('confidence', {}).get('detail', '')}"),
            ("EPSS 利用概率", f"{v.get('epss_score', 0):.1%} (超过 {v.get('epss_percentile', 0):.1%} 的漏洞)"),
            ("修复优先级", v.get("fix_priority", {}).get("label", "未知")),
            ("建议修复时间", v.get("fix_priority", {}).get("urgency", "未知")),
            ("修复版本", v.get("fixed_version", "未知")),
            ("CWE 类型", v.get("cwe_id", "未知")),
        ]

        detail_table = doc.add_table(rows=len(detail_data), cols=2, style='Light Grid Accent 1')
        for row_idx, (key, val) in enumerate(detail_data):
            _set_cell(detail_table, row_idx, 0, key)
            _set_cell(detail_table, row_idx, 1, val)

        if v.get("summary"):
            doc.add_paragraph(f'描述：{v["summary"]}')

        evidence_links = v.get("evidence_links", [])
        if evidence_links:
            doc.add_paragraph('证据来源：')
            for link in evidence_links:
                doc.add_paragraph(f'  - {link["label"]}: {link["url"]}')

        doc.add_paragraph('')

    # ===== 四、依赖清单 =====
    doc.add_heading('四、项目依赖清单', level=1)
    if dep_list:
        dep_table = doc.add_table(rows=1, cols=4, style='Light Grid Accent 1')
        _set_cell(dep_table, 0, 0, '包名')
        _set_cell(dep_table, 0, 1, '版本')
        _set_cell(dep_table, 0, 2, '生态')
        _set_cell(dep_table, 0, 3, '层级')
        for dep in dep_list[:50]:
            row = dep_table.add_row()
            row.cells[0].text = dep.get("package", "")
            row.cells[1].text = dep.get("version", "")
            row.cells[2].text = dep.get("ecosystem", "PyPI")
            row.cells[3].text = "间接（一层推断）" if dep.get("is_transitive") else "直接"
        if len(dep_list) > 50:
            doc.add_paragraph(f'... 共 {len(dep_list)} 个依赖，仅展示前 50 个')
    else:
        doc.add_paragraph('无依赖数据')

    doc.add_paragraph('')

    # ===== 五、修复建议 =====
    doc.add_heading('五、修复建议', level=1)

    priority_1 = [v for v in vulns if v.get("fix_priority", {}).get("priority") == 1]
    if priority_1:
        doc.add_paragraph(f'需立即修复 ({len(priority_1)} 个)：')
        for v in priority_1:
            doc.add_paragraph(
                f'  - {v["package"]} {v["version"]} → '
                f'升级到 {v.get("fixed_version", "最新版本")} '
                f'({v["cve_id"]}, CVSS {v["cvss_score"]})'
            )

    exploited = [v for v in vulns if v.get("is_actively_exploited")]
    if exploited:
        doc.add_paragraph(f'存在活跃利用风险 ({len(exploited)} 个)：')
        for v in exploited:
            doc.add_paragraph(f'  - {v["cve_id"]} — {v["package"]} {v["version"]}')

    doc.add_paragraph('通用安全建议：')
    doc.add_paragraph('  1. 定期运行依赖扫描')
    doc.add_paragraph('  2. 优先修复置信度"高"且修复优先级为"立即修复"的漏洞')
    doc.add_paragraph('  3. 关注 EPSS 高分的漏洞，这些是攻击者最可能利用的')
    doc.add_paragraph('  4. 引入新依赖前进行安全评估')

    # ===== 保存 =====
    os.makedirs("reports", exist_ok=True)
    report_path = f"reports/审计报告_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
    doc.save(report_path)
    return report_path


# ==========================================
# PDF 报告（精简版）
# ==========================================
def generate_pdf_report(scan_result: dict, sbom_data: dict) -> str:
    """生成 PDF 精简报告。"""
    from reportlab.pdfgen import canvas
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import cm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont

    # 注册内置中文字体（STSong-Light，跨平台可用，无需外部字体文件）
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))

    os.makedirs("reports", exist_ok=True)
    report_path = f"reports/审计报告_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"

    c = canvas.Canvas(report_path, pagesize=A4)
    width, height = A4
    y = height - 2 * cm
    summary = scan_result.get("summary", {})
    vulns = scan_result.get("vulns", [])

    c.setFont("Helvetica-Bold", 18)
    c.drawCentredString(width / 2, y, "Python Dependency Security Report")
    y -= 1.5 * cm

    c.setFont("Helvetica", 10)
    c.drawString(2 * cm, y, f"Scan Time: {summary.get('scan_time', 'N/A')}")
    y -= 0.5 * cm
    c.drawString(2 * cm, y, f"Dependencies: {summary.get('total_dependencies', 0)}")
    y -= 0.5 * cm
    c.drawString(2 * cm, y, f"Vulnerabilities: {summary.get('total_vulnerabilities', 0)}")
    y -= 0.5 * cm
    c.drawString(2 * cm, y, f"Critical/High: {summary.get('critical_count', 0)}")
    y -= 0.5 * cm
    c.drawString(2 * cm, y, f"High Confidence: {summary.get('high_confidence_count', 0)}")
    y -= 1 * cm

    c.setFont("Helvetica-Bold", 14)
    c.drawString(2 * cm, y, "Vulnerability List (Priority Order):")
    y -= 1 * cm
    c.setFont("Helvetica", 8)

    for v in vulns[:30]:
        fp = v.get("fix_priority", {}).get("priority", "?")
        line = f"[P{fp}] {v.get('cve_id', '')} | {v.get('package', '')} {v.get('version', '')} | {v.get('severity', '')} | CVSS:{v.get('cvss_score', 0)} | Confidence:{v.get('confidence', {}).get('level', '')} | Fix:{v.get('fixed_version', '')}"
        # 漏洞行含中文（severity/置信度），使用中文字体渲染
        c.setFont("STSong-Light", 8)
        c.drawString(2 * cm, y, line)
        y -= 0.4 * cm
        if y < 2 * cm:
            c.showPage()
            y = height - 2 * cm
            c.setFont("STSong-Light", 8)

    c.save()
    return report_path
